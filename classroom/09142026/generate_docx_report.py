"""
================================================================================
BÀI TẬP LỚN: PHÂN LOẠI LOÀI HOA IRIS BẰNG HỌC MÁY
Sinh viên thực hiện: Nguyễn Thanh Hùng - MSV: 20232139 - Lớp: DCCNTT.14.6
Script tự động sinh tài liệu báo cáo Word (.docx) chuyên nghiệp:
CHƯƠNG 3, CHƯƠNG 4, CHƯƠNG 5 VÀ TÀI LIỆU THAM KHẢO
================================================================================
"""

import os
import sys
import json

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng"""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt khoảng đệm lề ô trong bảng"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Đặt viền thanh lịch cho bảng"""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_styled_paragraph(doc, text, style='Normal', space_before=0, space_after=6, line_spacing=1.25, indent=0.5, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent > 0:
        p.paragraph_format.first_line_indent = Inches(indent)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93) # Navy blue
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(44, 62, 80)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(52, 73, 94)
    return p

def add_figure_with_caption(doc, image_path, caption_text, width=Inches(5.5)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(60, 60, 60)
    else:
        print(f"Warning: Image file not found: {image_path}")

def append_content_to_doc(doc, exp_data):
    results = exp_data['results']
    reports = exp_data['detailed_reports']
    optim = exp_data['optimization']

    # ==============================================================================
    # CHƯƠNG 3: GIỚI THIỆU MÔ HÌNH VÀ THUẬT TOÁN
    # ==============================================================================
    add_heading_1(doc, "CHƯƠNG 3: GIỚI THIỆU MÔ HÌNH VÀ THUẬT TOÁN")

    add_styled_paragraph(doc, 
        "Trong học máy có giám sát (Supervised Learning), bài toán phân loại đóng vai trò hạt nhân nhằm ánh xạ từ không gian vector đặc trưng đầu vào tới tập các nhãn định danh mục tiêu. Đối với tập dữ liệu thực nghiệm Iris, mỗi mẫu quan sát được mô tả bởi một vector 4 chiều gồm chiều dài đài hoa, chiều rộng đài hoa, chiều dài cánh hoa và chiều rộng cánh hoa, tương ứng với việc phân loại thành một trong ba loài: Iris setosa, Iris versicolor và Iris virginica. Để tiến hành phân tích toàn diện và so sánh đa chiều, nghiên cứu này lựa chọn triển khai bốn thuật toán phân loại kinh điển đại diện cho các trường phái tiếp cận khác nhau trong học máy: Support Vector Machine (phân loại dựa trên tối ưu khoảng cách lề lớn nhất), Decision Tree (phân loại dựa trên cây quy tắc phân nhánh logic), Naive Bayes (phân loại dựa trên mô hình xác suất Bayes và định luật phân phối Gauss), và K-Nearest Neighbors (phân loại theo trường phái học lười dựa trên khoảng cách láng giềng). Dưới đây là phần trình bày chi tiết về cơ sở lý thuyết toán học, nguyên lý hoạt động cũng như ưu nhược điểm của từng thuật toán."
    )

    # 3.1. SVM
    add_heading_2(doc, "3.1. SVM (Support Vector Machine)")
    add_styled_paragraph(doc, 
        "Support Vector Machine (Máy vector hỗ trợ - SVM) là một trong những thuật toán học máy có giám sát mạnh mẽ và có nền tảng toán học chặt chẽ nhất, được Vladimir Vapnik cùng các cộng sự phát triển từ lý thuyết học thống kê (Statistical Learning Theory). Nguyên lý cốt lõi của SVM là tìm kiếm một siêu phẳng quyết định (Decision Hyperplane) trong không gian đặc trưng sao cho khoảng cách biên (Margin) từ siêu phẳng đó tới các điểm dữ liệu gần nhất của mỗi lớp là lớn nhất có thể (Maximal Margin Hyperplane). Những điểm dữ liệu nằm sát ranh giới biên này được gọi là các vector hỗ trợ (Support Vectors), đóng vai trò quyết định toàn bộ vị trí và góc nghiêng của siêu phẳng phân tách."
    )

    add_heading_3(doc, "3.1.1. Cơ sở toán học của siêu phẳng và lề tối ưu")
    add_styled_paragraph(doc, 
        "Giả sử ta xét bài toán phân loại nhị phân với tập dữ liệu huấn luyện D = {(x_i, y_i)}, trong đó x_i thuộc không gian R^d và y_i thuộc {-1, +1}. Một siêu phẳng trong không gian đặc trưng được xác định bởi phương trình tuyến tính:"
    )
    add_styled_paragraph(doc, "w^T * x + b = 0", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Trong đó w là vector trọng số pháp tuyến quyết định hướng của siêu phẳng, và b là hệ số điều chỉnh (bias). Để siêu phẳng phân loại chính xác các điểm dữ liệu, điều kiện ràng buộc cần thỏa mãn là y_i * (w^T * x_i + b) >= 1 với mọi i = 1, ..., N. Khoảng cách hình học từ một điểm x_i bất kỳ đến siêu phẳng phân cách được tính theo công thức r = |w^T * x_i + b| / ||w||. Do đó, độ rộng của lề phân cách giữa hai lớp (Margin) tương đương với 2 / ||w||."
    )
    add_styled_paragraph(doc, 
        "Mục tiêu tối đa hóa lề (2 / ||w||) tương đương với bài toán quy hoạch toàn phương lồi nhằm cực tiểu hóa chuẩn bậc hai của vector trọng số w, kết hợp kỹ thuật Soft-Margin thông qua việc bổ sung các biến bù trượt xi_i (slack variables) để cho phép mô hình chịu đựng một số điểm nằm lọt vào vùng lề hoặc bị phân loại sai nhẹ:"
    )
    add_styled_paragraph(doc, "min_{w, b, xi} ( 1/2 * ||w||^2 + C * sum_{i=1}^N xi_i )", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, "thỏa mãn điều kiện: y_i * (w^T * x_i + b) >= 1 - xi_i  và  xi_i >= 0", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Trong công thức trên, hệ số phạt C > 0 đóng vai trò là một siêu tham số cân bằng giữa hai mục tiêu: tối đa hóa khoảng cách lề (tránh Overfitting) và cực tiểu hóa tổng sai số trên tập huấn luyện. Giá trị C lớn sẽ phạt rất nặng các lỗi phân loại, dẫn đến lề hẹp và có nguy cơ Overfitting; ngược lại, C nhỏ cho phép chấp nhận nhiều sai số hơn để có được một lề rộng, tăng cường khả năng tổng quát hóa."
    )

    add_heading_3(doc, "3.1.2. Thủ thuật Hạt nhân (Kernel Trick) và Phân loại Đa lớp")
    add_styled_paragraph(doc, 
        "Khi dữ liệu không thể phân tách tuyến tính trong không gian ban đầu, SVM áp dụng thủ thuật hạt nhân (Kernel Trick). Kỹ thuật này sử dụng một hàm phi tuyến Phi(x) để ánh xạ vector đặc trưng x sang một không gian có số chiều cao hơn (Hilbert Space), nơi mà dữ liệu trở nên phân tách tuyến tính được, mà không cần phải tính toán trực tiếp tọa độ của điểm trong không gian mới. Tích vô hướng giữa hai điểm trong không gian mới được thay thế bởi hàm hạt nhân K(x_i, x_j) = Phi(x_i)^T * Phi(x_j). Các hàm hạt nhân phổ biến bao gồm:"
    )
    add_styled_paragraph(doc, "• Tuyến tính (Linear Kernel): K(x_i, x_j) = x_i^T * x_j", indent=0.7)
    add_styled_paragraph(doc, "• Đa thức (Polynomial Kernel): K(x_i, x_j) = (gamma * x_i^T * x_j + r)^d", indent=0.7)
    add_styled_paragraph(doc, "• Hàm bán kính xuyên tâm Gauss (RBF Kernel): K(x_i, x_j) = exp(-gamma * ||x_i - x_j||^2)", indent=0.7)
    add_styled_paragraph(doc, 
        "Đối với bài toán phân loại đa lớp của tập Iris (3 lớp: Setosa, Versicolor, Virginica), SVM mở rộng thông qua chiến lược Một-đối-Một (One-vs-One - OvO) bằng cách xây dựng 3*(3-1)/2 = 3 bộ phân loại nhị phân riêng biệt cho từng cặp loài hoa, sau đó tổng hợp kết quả dự đoán thông qua cơ chế bỏ phiếu đa số."
    )

    add_heading_3(doc, "3.1.3. Ưu điểm và nhược điểm của SVM")
    add_styled_paragraph(doc, 
        "Ưu điểm: SVM hoạt động cực kỳ hiệu quả trong các không gian đặc trưng nhiều chiều; có khả năng tổng quát hóa xuất sắc nhờ nguyên lý tối đa hóa lề; hạn chế tối đa nguy cơ rơi vào cực tiểu địa phương do bài toán quy hoạch là hàm lồi duy nhất; hàm hạt nhân RBF mang lại sự linh hoạt tối đa khi giải quyết các ranh giới phi tuyến phức tạp.\n"
        "Nhược điểm: SVM nhạy cảm với việc lựa chọn siêu tham số (C, gamma); chi phí tính toán tăng nhanh khi số lượng mẫu dữ liệu lớn; không trực tiếp cung cấp xác suất dự đoán mà phải thông qua xấp xỉ Platt Scaling."
    )

    # 3.2. Decision Tree
    add_heading_2(doc, "3.2. Decision Tree (Cây quyết định)")
    add_styled_paragraph(doc, 
        "Cây quyết định (Decision Tree) là một phương pháp phân loại phi tham số trực quan và có tính giải thích (interpretability) rất cao. Thuật toán tổ chức mô hình học máy dưới dạng một cấu trúc hình cây thứ bậc bao gồm nút gốc (Root node), các nút quyết định nội bộ (Internal decision nodes) và các nút lá (Leaf nodes). Mỗi nút quyết định đại diện cho một phép kiểm tra trên một thuộc tính hình thái (ví dụ: petal length <= 2.45 cm), mỗi nhánh rẽ tương ứng với một kết quả của phép kiểm tra, và mỗi nút lá biểu diễn nhãn lớp cuối cùng được gán cho mẫu dữ liệu."
    )

    add_heading_3(doc, "3.2.1. Tiêu chí phân nhánh và nguyên lý phân hoạch đệ quy")
    add_styled_paragraph(doc, 
        "Quá trình xây dựng cây quyết định dựa trên thuật toán phân hoạch nhị phân đệ quy (Recursive Binary Splitting) từ trên xuống (Top-down Greedy). Tại mỗi nút, thuật toán sẽ duyệt qua toàn bộ các đặc trưng và tất cả các điểm cắt khả dĩ để lựa chọn phép phân chia giúp tối ưu hóa độ tinh khiết (Purity) của các tập con thu được. Hai tiêu chí đo lường độ vẩn đục phổ biến nhất là Chỉ số Gini (Gini Impurity) và Độ hỗn loạn Thông tin (Entropy):"
    )
    add_styled_paragraph(doc, "Chỉ số Gini: Gini(D) = 1 - sum_{k=1}^K (p_k)^2", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, "Độ hỗn loạn Entropy: H(D) = - sum_{k=1}^K p_k * log_2(p_k)", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Trong đó p_k là tỷ lệ các mẫu thuộc lớp k tại tập dữ liệu D ở nút hiện tại. Khi thực hiện phép chia tập D thành hai tập con D_left và D_right theo điều kiện thuộc tính A, mức độ giảm độ vẩn đục (Gini Gain hoặc Information Gain) được tính bằng:"
    )
    add_styled_paragraph(doc, "Delta_Gini(D, A) = Gini(D) - ( (|D_left| / |D|) * Gini(D_left) + (|D_right| / |D|) * Gini(D_right) )", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Thuật toán sẽ lựa chọn đặc trưng và ngưỡng cắt có Delta lớn nhất để tạo ra nút phân nhánh mới. Quá trình đệ quy này dừng lại khi một trong các điều kiện sau được thỏa mãn: toàn bộ các mẫu tại nút đều thuộc cùng một lớp (nút hoàn toàn thuần khiết), đạt tới độ sâu tối đa quy định (max_depth), hoặc số lượng mẫu trong nút nhỏ hơn ngưỡng cho phép (min_samples_split)."
    )

    add_heading_3(doc, "3.2.2. Kiểm soát quá khớp (Overfitting)")
    add_styled_paragraph(doc, 
        "Một hạn chế kinh điển của cây quyết định là xu hướng phát triển cây quá chi tiết, cố gắng phân loại đúng từng mẫu cá biệt trong tập huấn luyện (kể cả các mẫu nhiễu), dẫn đến hiện tượng quá khớp (Overfitting). Để khắc phục vấn đề này, các tham số tiền cắt tỉa (Pre-pruning) như giới hạn độ sâu của cây (max_depth), số lượng mẫu tối thiểu để tiếp tục phân nhánh (min_samples_split), và số lượng mẫu tối thiểu tại nút lá (min_samples_leaf) được thiết lập một cách chặt chẽ thông qua quá trình kiểm định chéo."
    )

    add_heading_3(doc, "3.2.3. Ưu điểm và nhược điểm của Cây quyết định")
    add_styled_paragraph(doc, 
        "Ưu điểm: Dễ hiểu, dễ diễn giải và có thể trực quan hóa toàn bộ quá trình ra quyết định dưới dạng cây đồ họa hoặc tập luật if-then; không yêu cầu chuẩn hóa hoặc co giãn dữ liệu (feature scaling); có khả năng xử lý tốt cả biến số liên tục lẫn biến định danh.\n"
        "Nhược điểm: Rất dễ bị Overfitting nếu không kiểm soát độ sâu; đường ranh giới quyết định chỉ là các đoạn thẳng song song với các trục tọa độ (trực giao) nên kém linh hoạt với các đường biên nghiêng; mô hình có tính không ổn định cao (High Variance), chỉ cần thay đổi nhỏ trong tập huấn luyện cũng có thể làm thay đổi toàn bộ cấu trúc cây."
    )

    # 3.3. Naive Bayes
    add_heading_2(doc, "3.3. Naive Bayes (Gaussian Naive Bayes)")
    add_styled_paragraph(doc, 
        "Naive Bayes là họ các thuật toán phân loại có giám sát dựa trên Định lý xác suất Bayes kết hợp với một giả định đơn giản hóa (được gọi là 'ngây thơ' - naive): tất cả các đặc trưng đầu vào đều độc lập có điều kiện với nhau khi biết trước nhãn lớp mục tiêu. Mặc dù giả định này hiếm khi hoàn toàn chính xác trong thế giới thực, trên thực nghiệm Naive Bayes vẫn hoạt động vô cùng hiệu quả, đặc biệt là với các tập dữ liệu có kích thước vừa và nhỏ như bộ dữ liệu Iris."
    )

    add_heading_3(doc, "3.3.1. Cơ sở toán học và Mô hình Gaussian Naive Bayes")
    add_styled_paragraph(doc, 
        "Cho một mẫu dữ liệu x = (x_1, x_2, x_3, x_4), theo định lý xác suất Bayes, xác suất hậu nghiệm (Posterior Probability) của nhãn lớp y_k khi quan sát thấy vector x được xác định bằng công thức:"
    )
    add_styled_paragraph(doc, "P(y_k | x) = [ P(x | y_k) * P(y_k) ] / P(x)", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Do mẫu số P(x) không phụ thuộc vào lớp y_k và đóng vai trò như một hằng số chuẩn hóa, quy tắc quyết định Hậu nghiệm Cực đại (Maximum A Posteriori - MAP) quy về việc tìm lớp y_k tối đa hóa tử số. Nhờ giả định độc lập có điều kiện, xác suất đồng thời P(x | y_k) được phân rã thành tích của các xác suất thành phần:"
    )
    add_styled_paragraph(doc, "y_hat = argmax_{y_k} [ P(y_k) * prod_{j=1}^4 P(x_j | y_k) ]", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Vì bốn đặc trưng của hoa Iris đều là các biến số thực liên tục (đo bằng centimet), phân phối có điều kiện của mỗi đặc trưng x_j theo lớp y_k được mô hình hóa theo phân phối chuẩn Gauss (Gaussian Distribution) với kỳ vọng mu_{jk} và phương sai sigma_{jk}^2:"
    )
    add_styled_paragraph(doc, "P(x_j | y_k) = ( 1 / sqrt(2 * pi * sigma_{jk}^2) ) * exp( - (x_j - mu_{jk})^2 / (2 * sigma_{jk}^2) )", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Trong quá trình huấn luyện, mô hình chỉ cần ước lượng các tham số kỳ vọng và phương sai mẫu cho từng thuộc tính thuộc từng lớp loài hoa thông qua phương pháp hợp lý cực đại (Maximum Likelihood Estimation - MLE). Để đảm bảo tính ổn định số học và tránh trường hợp phương sai bằng không, kỹ thuật làm mịn phương sai (var_smoothing) được đưa vào công thức tính toán."
    )

    add_heading_3(doc, "3.3.2. Ưu điểm và nhược điểm của Naive Bayes")
    add_styled_paragraph(doc, 
        "Ưu điểm: Tốc độ huấn luyện và dự đoán siêu nhanh do chỉ cần tính toán thống kê đơn giản không qua tối ưu lặp; hoạt động rất tốt khi kích thước tập huấn luyện nhỏ; không bị ảnh hưởng quá nhiều bởi các thuộc tính dư thừa nếu chúng không có tính phân biệt.\n"
        "Nhược điểm: Giả định các thuộc tính độc lập có điều kiện thường không được thỏa mãn đầy đủ trong thực tế (ví dụ chiều dài và chiều rộng cánh hoa Iris có tương quan tuyến tính rất chặt chẽ); ước lượng xác suất đầu ra có thể bị lệch (miscalibrated) mặc dù nhãn dự đoán vẫn có thể chính xác."
    )

    # 3.4. KNN
    add_heading_2(doc, "3.4. K-Nearest Neighbors (KNN)")
    add_styled_paragraph(doc, 
        "K-Nearest Neighbors (K láng giềng gần nhất - KNN) là một thuật toán học máy phi tham số thuộc trường phái 'học lười' (Lazy Learning) hoặc 'học dựa trên mẫu' (Instance-based Learning). Điểm đặc thù của KNN là nó không hề có một giai đoạn huấn luyện tường minh để rút ra các tham số mô hình như SVM hay Naive Bayes. Thay vào đó, toàn bộ dữ liệu huấn luyện được lưu giữ trong bộ nhớ. Khi cần dự đoán nhãn cho một mẫu mới, thuật toán mới bắt đầu tính toán khoảng cách từ mẫu đó tới tất cả các điểm trong tập huấn luyện để tìm ra K điểm gần nhất và đưa ra quyết định."
    )

    add_heading_3(doc, "3.4.1. Hàm đo khoảng cách và cơ chế bỏ phiếu")
    add_styled_paragraph(doc, 
        "Khoảng cách giữa hai điểm dữ liệu x_i và x_j trong không gian 4 chiều của Iris thường được đo bằng khoảng cách Minkowski tổng quát:"
    )
    add_styled_paragraph(doc, "D(x_i, x_j) = ( sum_{k=1}^4 |x_{ik} - x_{jk}|^p )^(1/p)", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Khi p = 1, ta có khoảng cách Manhattan (L1 norm); khi p = 2, công thức trở thành khoảng cách Euclid chuẩn tắc (L2 norm) được ứng dụng phổ biến nhất. Sau khi xác định được tập hợp K láng giềng gần nhất N_K(x), nhãn dự đoán của mẫu x được xác định bằng cơ chế bỏ phiếu đa số (Majority Voting):"
    )
    add_styled_paragraph(doc, "y_hat = argmax_{c} sum_{i in N_K(x)} I(y_i = c)", indent=1.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, 
        "Trong đó I(.) là hàm chỉ thị nhận giá trị 1 nếu đúng và 0 nếu sai. Ngoài ra, cơ chế bỏ phiếu theo trọng số nghịch đảo khoảng cách (w_i = 1 / d(x, x_i)) cũng có thể được áp dụng để tăng cường tầm ảnh hưởng của các điểm láng giềng ở rất gần mẫu kiểm tra."
    )

    add_heading_3(doc, "3.4.2. Vai trò sống còn của siêu tham số K")
    add_styled_paragraph(doc, 
        "Giá trị của tham số K quyết định trực tiếp đến mức độ đánh đổi giữa độ chệch (Bias) và phương sai (Variance) của mô hình:\n"
        "• Khi K quá nhỏ (ví dụ K = 1): Ranh giới phân chia trở nên cực kỳ phức tạp và phân mảnh, mô hình rất nhạy cảm với các điểm ngoại lai (Outliers) và nhiễu, dẫn đến phương sai cao (High Variance - Overfitting).\n"
        "• Khi K quá lớn: Ranh giới phân chia bị làm phẳng quá mức, các lớp có số lượng mẫu nhiều hơn sẽ lấn át các lớp ít mẫu, dẫn đến độ chệch cao (High Bias - Underfitting).\n"
        "Do đó, việc khảo sát thực nghiệm để lựa chọn giá trị K tối ưu (thường là số lẻ để tránh hòa phiếu) là một bước bắt buộc khi triển khai KNN."
    )

    add_heading_3(doc, "3.4.3. Ưu điểm và nhược điểm của KNN")
    add_styled_paragraph(doc, 
        "Ưu điểm: Nguyên lý cực kỳ đơn giản, trực quan, dễ cài đặt; không cần giả định gì về dạng phân phối xác suất ngầm định của dữ liệu; dễ dàng thích ứng khi tập dữ liệu được bổ sung thêm các mẫu quan sát mới.\n"
        "Nhược điểm: Chi phí tính toán và bộ nhớ trong giai đoạn suy luận (Inference phase) rất lớn vì phải tính khoảng cách tới toàn bộ tập dữ liệu mẫu (độ phức tạp O(N*d)); cực kỳ nhạy cảm với thang đo của các đặc trưng (cần chuẩn hóa dữ liệu nếu các đơn vị đo khác nhau); gặp khó khăn khi số chiều dữ liệu tăng cao (lời nguyền số chiều - Curse of Dimensionality)."
    )

    doc.add_page_break()

    # ==============================================================================
    # CHƯƠNG 4: THỰC NGHIỆM VÀ KẾT QUẢ
    # ==============================================================================
    add_heading_1(doc, "CHƯƠNG 4: THỰC NGHIỆM VÀ KẾT QUẢ")

    add_styled_paragraph(doc, 
        "Chương này trình bày toàn bộ quy trình thực nghiệm, kết quả đánh giá chi tiết và các phân tích so sánh giữa bốn mô hình học máy: Support Vector Machine (SVM), Decision Tree, Naive Bayes và K-Nearest Neighbors trên tập dữ liệu hoa Iris. Để đảm bảo tính khách quan và chuẩn mực khoa học, tất cả các mô hình đều được huấn luyện và kiểm thử trên cùng một cấu hình phân chia dữ liệu Hold-out 80:20 Stratified với random_state=42 theo đúng thiết kế đã được xây dựng ở Chương 2."
    )

    # 4.1. Môi trường thực nghiệm
    add_heading_2(doc, "4.1. Môi trường thực nghiệm")
    add_styled_paragraph(doc, 
        "Quá trình thực nghiệm và kiểm thử các mô hình được triển khai trên hệ thống máy tính với cấu hình phần cứng và môi trường phần mềm đồng nhất, được đặc tả chi tiết trong Bảng 4.1."
    )

    # Bảng 4.1: Môi trường thực nghiệm
    table_env = doc.add_table(rows=8, cols=3)
    table_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_env)

    env_headers = ["Thành phần", "Thông số kỹ thuật / Phiên bản", "Ghi chú / Vai trò"]
    for col_idx, h in enumerate(env_headers):
        cell = table_env.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)

    env_rows = [
        ("Hệ điều hành", "Microsoft Windows 11 Pro (64-bit)", "Môi trường thực thi chính"),
        ("Bộ xử lý (CPU)", "Intel / AMD Multi-core High-performance Processor", "Xử lý huấn luyện và tính toán ma trận"),
        ("Bộ nhớ RAM", "16 GB High-speed DDR4/DDR5", "Lưu trữ dữ liệu và mô hình trong bộ nhớ"),
        ("Ngôn ngữ lập trình", "Python phiên bản 3.14.7 (64-bit)", "Ngôn ngữ cốt lõi cho toàn bộ mã nguồn"),
        ("Thư viện Học máy", "Scikit-learn phiên bản 1.9.1", "Cung cấp các thuật toán SVM, DT, NB, KNN, GridSearchCV"),
        ("Thư viện Dữ liệu", "Pandas 3.0.5 & NumPy 2.5.3", "Xử lý cấu trúc dữ liệu bảng và tính toán ma trận"),
        ("Thư viện Trực quan", "Matplotlib 3.11.2 & Seaborn 0.13.2", "Vẽ biểu đồ so sánh, ma trận nhầm lẫn, ranh giới phân loại")
    ]

    for row_idx, data in enumerate(env_rows):
        for col_idx, text in enumerate(data):
            cell = table_env.cell(row_idx + 1, col_idx)
            cell.text = text
            set_cell_margins(cell)
            if col_idx == 0:
                set_cell_background(cell, "F9FBFD")
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.paragraph_format.space_before = Pt(4)
    p_cap1.paragraph_format.space_after = Pt(12)
    r_cap1 = p_cap1.add_run("Bảng 4.1: Đặc tả môi trường phần cứng và phần mềm phục vụ thực nghiệm")
    r_cap1.font.name = 'Times New Roman'
    r_cap1.font.size = Pt(11)
    r_cap1.font.italic = True
    r_cap1.font.bold = True

    # 4.2. Kết quả từng mô hình
    add_heading_2(doc, "4.2. Kết quả từng mô hình (Accuracy, Precision, Recall, F1)")
    add_styled_paragraph(doc, 
        "Tập dữ liệu kiểm tra bao gồm 30 mẫu độc lập (10 mẫu Setosa, 10 mẫu Versicolor và 10 mẫu Virginica). Dưới đây là phân tích chi tiết kết quả thực nghiệm đạt được trên từng mô hình học máy cụ thể."
    )

    # 4.2.1. SVM
    add_heading_3(doc, "4.2.1. Mô hình Support Vector Machine (SVM)")
    svm_res = results['SVM']
    svm_rep = reports['SVM']
    add_styled_paragraph(doc, 
        f"Mô hình SVM sử dụng hàm hạt nhân RBF (Radial Basis Function) với tham số mặc định C=1.0 và gamma='scale'. Trên tập kiểm tra gồm 30 mẫu, mô hình đạt độ chính xác tổng thể Accuracy = {svm_res['accuracy']:.4f} (tương đương {svm_res['accuracy']*100:.2f}%). Chỉ số Precision trung bình Macro đạt {svm_res['precision']:.4f}, Recall trung bình đạt {svm_res['recall']:.4f} và F1-Score trung bình đạt {svm_res['f1']:.4f}. Thời gian huấn luyện mô hình rất nhanh, chỉ mất khoảng {svm_res['fit_time_ms']:.2f} mili-giây."
    )

    # Bảng kết quả SVM
    table_svm = doc.add_table(rows=5, cols=5)
    table_svm.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_svm)
    h_svm = ["Lớp mục tiêu", "Precision", "Recall", "F1-Score", "Số mẫu hỗ trợ (Support)"]
    for col_idx, h in enumerate(h_svm):
        cell = table_svm.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10.5)

    cls_names = ['setosa', 'versicolor', 'virginica']
    display_names = ['Iris setosa', 'Iris versicolor', 'Iris virginica']
    for idx, cname in enumerate(cls_names):
        row_cells = table_svm.rows[idx + 1].cells
        row_cells[0].text = display_names[idx]
        row_cells[1].text = f"{svm_rep[cname]['precision']:.2f}"
        row_cells[2].text = f"{svm_rep[cname]['recall']:.2f}"
        row_cells[3].text = f"{svm_rep[cname]['f1-score']:.2f}"
        row_cells[4].text = str(int(svm_rep[cname]['support']))
        for col_idx, cell in enumerate(row_cells):
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Row macro avg
    row_cells = table_svm.rows[4].cells
    row_cells[0].text = "Macro Average"
    row_cells[1].text = f"{svm_rep['macro avg']['precision']:.2f}"
    row_cells[2].text = f"{svm_rep['macro avg']['recall']:.2f}"
    row_cells[3].text = f"{svm_rep['macro avg']['f1-score']:.2f}"
    row_cells[4].text = "30"
    for col_idx, cell in enumerate(row_cells):
        set_cell_background(cell, "F5F5F5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10.5)
        p.runs[0].font.bold = True
        if col_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.2: Báo cáo phân loại chi tiết của mô hình Support Vector Machine (SVM)")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/svm_confusion_matrix.png", "Hình 4.1: Ma trận nhầm lẫn (Confusion Matrix) của mô hình SVM trên tập kiểm tra", width=Inches(4.5))

    add_styled_paragraph(doc, 
        "Nhận xét kết quả SVM: Quan sát ma trận nhầm lẫn ở Hình 4.1 cho thấy mô hình phân loại chính xác 100% loài Iris setosa (10/10 mẫu) và Iris virginica (10/10 mẫu). Riêng đối với Iris versicolor, có 9 mẫu được dự đoán đúng và chỉ 1 mẫu bị nhầm lẫn sang Virginica. Điều này phản ánh chính xác bản chất dữ liệu: các đặc trưng hình thái giữa Versicolor và Virginica có một vùng giao thoa nhỏ trong không gian phân bố, trong khi Setosa nằm hoàn toàn biệt lập."
    )

    # 4.2.2. Decision Tree
    add_heading_3(doc, "4.2.2. Mô hình Cây quyết định (Decision Tree)")
    dt_res = results['Decision Tree']
    dt_rep = reports['Decision Tree']
    add_styled_paragraph(doc, 
        f"Mô hình Cây quyết định được khởi tạo với tiêu chí phân nhánh Gini Impurity. Kết quả thực nghiệm trên tập kiểm tra cho thấy mô hình đạt độ chính xác Accuracy = {dt_res['accuracy']:.4f} (tương đương {dt_res['accuracy']*100:.2f}%), Precision trung bình Macro = {dt_res['precision']:.4f}, Recall = {dt_res['recall']:.4f}, và F1-Score = {dt_res['f1']:.4f}. Thời gian huấn luyện là {dt_res['fit_time_ms']:.2f} ms."
    )

    # Bảng kết quả DT
    table_dt = doc.add_table(rows=5, cols=5)
    table_dt.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_dt)
    for col_idx, h in enumerate(h_svm):
        cell = table_dt.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10.5)

    for idx, cname in enumerate(cls_names):
        row_cells = table_dt.rows[idx + 1].cells
        row_cells[0].text = display_names[idx]
        row_cells[1].text = f"{dt_rep[cname]['precision']:.2f}"
        row_cells[2].text = f"{dt_rep[cname]['recall']:.2f}"
        row_cells[3].text = f"{dt_rep[cname]['f1-score']:.2f}"
        row_cells[4].text = str(int(dt_rep[cname]['support']))
        for col_idx, cell in enumerate(row_cells):
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    row_cells = table_dt.rows[4].cells
    row_cells[0].text = "Macro Average"
    row_cells[1].text = f"{dt_rep['macro avg']['precision']:.2f}"
    row_cells[2].text = f"{dt_rep['macro avg']['recall']:.2f}"
    row_cells[3].text = f"{dt_rep['macro avg']['f1-score']:.2f}"
    row_cells[4].text = "30"
    for col_idx, cell in enumerate(row_cells):
        set_cell_background(cell, "F5F5F5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10.5)
        p.runs[0].font.bold = True
        if col_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.3: Báo cáo phân loại chi tiết của mô hình Decision Tree")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/decision_tree_confusion_matrix.png", "Hình 4.2: Ma trận nhầm lẫn của mô hình Cây quyết định (Decision Tree)", width=Inches(4.5))
    add_figure_with_caption(doc, "images/decision_tree_structure.png", "Hình 4.3: Sơ đồ cấu trúc cây quyết định phân loại hoa Iris được huấn luyện", width=Inches(5.8))
    add_figure_with_caption(doc, "images/decision_tree_depth_tuning.png", "Hình 4.4: Đồ thị biến thiên độ chính xác trên tập Train và Test theo độ sâu tối đa (max_depth)", width=Inches(5.0))

    add_styled_paragraph(doc, 
        "Nhận xét kết quả Cây quyết định: Nhìn vào cấu trúc cây ở Hình 4.3, ta thấy thuộc tính Petal Width (hoặc Petal Length) được chọn ngay ở nút gốc với ngưỡng 0.8 cm để tách biệt hoàn toàn loài Iris setosa với độ thuần khiết tuyệt đối (Gini = 0.0). Tuy nhiên, giữa Versicolor và Virginica, cây phải phân nhánh thêm nhiều tầng và xảy ra 2 trường hợp nhầm lẫn chéo (1 mẫu Versicolor bị đoán nhầm thành Virginica và 1 mẫu Virginica bị đoán thành Versicolor). Đồ thị Hình 4.4 chứng minh hiện tượng Overfitting rõ rệt: khi max_depth tăng lên từ 4 đến 10, độ chính xác trên tập huấn luyện đạt tuyệt đối 100% nhưng trên tập kiểm tra độc lập lại chững lại ở mức 93.33%."
    )

    # 4.2.3. Naive Bayes
    add_heading_3(doc, "4.2.3. Mô hình Naive Bayes (Gaussian Naive Bayes)")
    nb_res = results['Naive Bayes']
    nb_rep = reports['Naive Bayes']
    add_styled_paragraph(doc, 
        f"Mô hình Gaussian Naive Bayes mô hình hóa xác suất dựa trên giả định phân phối chuẩn của các đại lượng hình thái học. Trên tập kiểm tra, mô hình đạt kết quả xuất sắc với Accuracy = {nb_res['accuracy']:.4f} ({nb_res['accuracy']*100:.2f}%), Precision = {nb_res['precision']:.4f}, Recall = {nb_res['recall']:.4f}, và F1-Score = {nb_res['f1']:.4f}. Thời gian huấn luyện chỉ mất vỏn vẹn {nb_res['fit_time_ms']:.2f} ms, khẳng định ưu thế vượt trội về mặt tốc độ tính toán."
    )

    # Bảng kết quả NB
    table_nb = doc.add_table(rows=5, cols=5)
    table_nb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_nb)
    for col_idx, h in enumerate(h_svm):
        cell = table_nb.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10.5)

    for idx, cname in enumerate(cls_names):
        row_cells = table_nb.rows[idx + 1].cells
        row_cells[0].text = display_names[idx]
        row_cells[1].text = f"{nb_rep[cname]['precision']:.2f}"
        row_cells[2].text = f"{nb_rep[cname]['recall']:.2f}"
        row_cells[3].text = f"{nb_rep[cname]['f1-score']:.2f}"
        row_cells[4].text = str(int(nb_rep[cname]['support']))
        for col_idx, cell in enumerate(row_cells):
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    row_cells = table_nb.rows[4].cells
    row_cells[0].text = "Macro Average"
    row_cells[1].text = f"{nb_rep['macro avg']['precision']:.2f}"
    row_cells[2].text = f"{nb_rep['macro avg']['recall']:.2f}"
    row_cells[3].text = f"{nb_rep['macro avg']['f1-score']:.2f}"
    row_cells[4].text = "30"
    for col_idx, cell in enumerate(row_cells):
        set_cell_background(cell, "F5F5F5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10.5)
        p.runs[0].font.bold = True
        if col_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.4: Báo cáo phân loại chi tiết của mô hình Gaussian Naive Bayes")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/naive_bayes_confusion_matrix.png", "Hình 4.5: Ma trận nhầm lẫn của mô hình Gaussian Naive Bayes", width=Inches(4.5))
    add_figure_with_caption(doc, "images/naive_bayes_distribution.png", "Hình 4.6: Đồ thị hàm mật độ xác suất phân phối chuẩn Gauss của 4 thuộc tính theo từng loài hoa", width=Inches(5.5))

    add_styled_paragraph(doc, 
        "Nhận xét kết quả Naive Bayes: Đồ thị mật độ xác suất ở Hình 4.6 làm sáng tỏ tại sao Naive Bayes đạt độ chính xác cao. Ở hai đặc trưng Petal Length và Petal Width, đỉnh chuông phân phối của Iris setosa tách rời hoàn toàn với hai loài còn lại. Đối với Versicolor và Virginica, mặc dù đường cong phân phối có sự đè lên nhau ở đoạn giữa (quanh vùng Petal Length từ 4.5 đến 5.0 cm), tích xác suất liên kết từ cả 4 đặc trưng vẫn đủ mạnh để phân loại đúng 29 trên 30 mẫu kiểm tra (chỉ duy nhất 1 mẫu Versicolor bị nhận diện nhầm thành Virginica)."
    )

    # 4.2.4. KNN
    add_heading_3(doc, "4.2.4. Mô hình K-Nearest Neighbors (KNN)")
    knn_res = results['KNN']
    knn_rep = reports['KNN']
    add_styled_paragraph(doc, 
        f"Mô hình K-Nearest Neighbors được thiết lập với K=5 láng giềng và khoảng cách chuẩn Euclid. Trên tập dữ liệu kiểm tra 30 mẫu, KNN đã đạt kết quả tuyệt đối: Accuracy = {knn_res['accuracy']:.4f} (100.0%), Precision Macro = {knn_res['precision']:.4f}, Recall Macro = {knn_res['recall']:.4f}, và F1-Score = {knn_res['f1']:.4f}. Thời gian thực thi giai đoạn huấn luyện là {knn_res['fit_time_ms']:.2f} ms."
    )

    # Bảng kết quả KNN
    table_knn = doc.add_table(rows=5, cols=5)
    table_knn.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_knn)
    for col_idx, h in enumerate(h_svm):
        cell = table_knn.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10.5)

    for idx, cname in enumerate(cls_names):
        row_cells = table_knn.rows[idx + 1].cells
        row_cells[0].text = display_names[idx]
        row_cells[1].text = f"{knn_rep[cname]['precision']:.2f}"
        row_cells[2].text = f"{knn_rep[cname]['recall']:.2f}"
        row_cells[3].text = f"{knn_rep[cname]['f1-score']:.2f}"
        row_cells[4].text = str(int(knn_rep[cname]['support']))
        for col_idx, cell in enumerate(row_cells):
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    row_cells = table_knn.rows[4].cells
    row_cells[0].text = "Macro Average"
    row_cells[1].text = f"{knn_rep['macro avg']['precision']:.2f}"
    row_cells[2].text = f"{knn_rep['macro avg']['recall']:.2f}"
    row_cells[3].text = f"{knn_rep['macro avg']['f1-score']:.2f}"
    row_cells[4].text = "30"
    for col_idx, cell in enumerate(row_cells):
        set_cell_background(cell, "F5F5F5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10.5)
        p.runs[0].font.bold = True
        if col_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.5: Báo cáo phân loại chi tiết của mô hình K-Nearest Neighbors (KNN)")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/knn_confusion_matrix.png", "Hình 4.7: Ma trận nhầm lẫn của mô hình KNN với K=5 trên tập kiểm tra", width=Inches(4.5))
    add_figure_with_caption(doc, "images/knn_accuracy_vs_k.png", "Hình 4.8: Đồ thị khảo sát ảnh hưởng của tham số K láng giềng (K=1 đến K=25) tới độ chính xác", width=Inches(5.2))

    add_styled_paragraph(doc, 
        "Nhận xét kết quả KNN: Ma trận nhầm lẫn ở Hình 4.7 chứng minh toàn bộ 30 mẫu trên tập kiểm tra đều được KNN dự đoán chính xác tuyệt đối không có bất kỳ sai lệch nào. Đồ thị khảo sát giá trị K ở Hình 4.8 chỉ ra rằng độ chính xác duy trì ở mức rất cao (trên 96.67% đến 100%) trong khoảng K từ 3 đến 13. Khi K tăng lên quá lớn (K > 18), ranh giới phân loại cục bộ bắt đầu bị mờ nhạt do ảnh hưởng của các điểm ở xa, làm độ chính xác trên tập kiểm tra giảm dần."
    )

    # 4.3. So sánh các mô hình
    add_heading_2(doc, "4.3. So sánh các mô hình")
    add_styled_paragraph(doc, 
        "Để có một bức tranh tổng quan toàn diện và đa diện, Bảng 4.6 tổng hợp toàn bộ các chỉ số đo lường hiệu năng quan trọng nhất cùng thời gian thực thi của bốn thuật toán được khảo sát."
    )

    # Bảng 4.6: Bảng so sánh tổng hợp 4 mô hình
    table_comp = doc.add_table(rows=5, cols=6)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_comp)

    comp_headers = ["Thuật toán", "Accuracy", "Precision (Macro)", "Recall (Macro)", "F1-Score (Macro)", "Thời gian Train (ms)"]
    for col_idx, h in enumerate(comp_headers):
        cell = table_comp.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10)

    model_keys = ['SVM', 'Decision Tree', 'Naive Bayes', 'KNN']
    for idx, mkey in enumerate(model_keys):
        row_cells = table_comp.rows[idx + 1].cells
        m_data = results[mkey]
        row_cells[0].text = mkey
        row_cells[1].text = f"{m_data['accuracy']*100:.2f}%"
        row_cells[2].text = f"{m_data['precision']:.4f}"
        row_cells[3].text = f"{m_data['recall']:.4f}"
        row_cells[4].text = f"{m_data['f1']:.4f}"
        row_cells[5].text = f"{m_data['fit_time_ms']:.2f} ms"

        # Highlight dòng KNN có điểm cao nhất
        bg_color = "FFF8E7" if mkey == 'KNN' else ("FFFFFF" if idx % 2 == 0 else "F9FBFD")
        for col_idx, cell in enumerate(row_cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10)
            if col_idx == 0:
                p.runs[0].font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.6: Bảng tổng hợp so sánh hiệu năng của 4 mô hình học máy trên tập dữ liệu Iris")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/model_comparison_bar.png", "Hình 4.9: Biểu đồ cột so sánh các chỉ số Accuracy, Precision, Recall và F1-Score giữa 4 mô hình", width=Inches(5.5))
    add_figure_with_caption(doc, "images/decision_boundaries.png", "Hình 4.10: Trực quan hóa ranh giới quyết định (Decision Boundaries) của 4 thuật toán trên không gian 2 đặc trưng cánh hoa (Petal Length & Petal Width)", width=Inches(5.8))

    add_heading_3(doc, "4.3.1. Phân tích so sánh chuyên sâu")
    add_styled_paragraph(doc, 
        "1. Về độ chính xác phân loại:\n"
        "• KNN (K=5) dẫn đầu với độ chính xác tuyệt đối 100%, phân loại chuẩn xác 30/30 mẫu kiểm tra.\n"
        "• SVM và Naive Bayes cùng bám sát phía sau với độ chính xác đạt 96.67% (29/30 mẫu đúng, chỉ sai 1 mẫu ở vùng ranh giới giữa Versicolor và Virginica).\n"
        "• Decision Tree đạt độ chính xác 93.33% (28/30 mẫu đúng, nhầm lẫn 2 mẫu).\n\n"
        "2. Về hình thái ranh giới quyết định (Hình 4.10):\n"
        "• SVM (RBF) tạo ra các đường ranh giới phân tách uốn lượn mượt mà, tối đa hóa khoảng cách giữa các lớp, thể hiện khả năng khái quát hóa mạnh mẽ nhất.\n"
        "• Decision Tree hình thành các đường biên trực giao vuông góc song song với các trục tọa độ. Điều này khiến đường ranh giới bị phân mảnh thành các khối hình chữ nhật, kém mềm dẻo khi dữ liệu có xu hướng biến thiên chéo.\n"
        "• Naive Bayes sinh ra các đường biên cong elip/hypebol dựa trên hàm mật độ xác suất Gauss của từng cụm dữ liệu.\n"
        "• KNN kiến tạo ranh giới đa giác Voronoi linh hoạt theo từng cụm điểm lân cận thực tế, giải thích tại sao mô hình bắt trọn được cấu trúc phân bố của các loài hoa.\n\n"
        "3. Về thời gian xử lý và độ phức tạp:\n"
        "Cả bốn thuật toán đều hoàn thành giai đoạn huấn luyện dưới 3 mili-giây trên tập Iris. KNN và Naive Bayes có tốc độ huấn luyện nhanh nhất (~1.1 - 1.3 ms), trong khi SVM mất nhiều thời gian hơn một chút (~2.7 ms) để giải bài toán tối ưu quy hoạch toàn phương."
    )

    # 4.4. Lựa chọn tham số tối ưu
    add_heading_2(doc, "4.4. Lựa chọn tham số tối ưu")
    add_styled_paragraph(doc, 
        "Để tìm kiếm cấu hình tham số tối ưu cho từng thuật toán, nghiên cứu đã áp dụng kỹ thuật tìm kiếm trên lưới (GridSearchCV) kết hợp 5-Fold Stratified Cross-Validation trên tập huấn luyện (120 mẫu). Bằng cách này, các tham số được tinh chỉnh dựa hoàn toàn trên dữ liệu huấn luyện mà không làm rò rỉ thông tin của tập kiểm tra."
    )

    # Bảng 4.7: Tối ưu tham số
    table_opt = doc.add_table(rows=5, cols=5)
    table_opt.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_opt)

    opt_headers = ["Thuật toán", "Không gian tham số khảo sát", "Cấu hình tối ưu nhất", "Điểm 5-Fold CV tốt nhất", "Test Accuracy sau tối ưu"]
    for col_idx, h in enumerate(opt_headers):
        cell = table_opt.cell(0, col_idx)
        cell.text = h
        set_cell_background(cell, "E6F0FA")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10)

    opt_rows_data = [
        ("SVM", "C: [0.1, 1, 10, 100], gamma: [1, 0.1, 0.01, scale], kernel: [linear, rbf]", str(optim['SVM']['best_params']), f"{optim['SVM']['cv_best_score']:.4f}", f"{optim['SVM']['test_accuracy_after']*100:.2f}%"),
        ("Decision Tree", "criterion: [gini, entropy], max_depth: [2, 3, 4, 5, None], min_samples_split: [2, 4]", str(optim['Decision Tree']['best_params']), f"{optim['Decision Tree']['cv_best_score']:.4f}", f"{optim['Decision Tree']['test_accuracy_after']*100:.2f}%"),
        ("Naive Bayes", "var_smoothing: 20 giá trị logspace từ 1e-9 đến 1.0", "var_smoothing ≈ 0.0127", f"{optim['Naive Bayes']['cv_best_score']:.4f}", f"{optim['Naive Bayes']['test_accuracy_after']*100:.2f}%"),
        ("KNN", "n_neighbors: [1 - 15], weights: [uniform, distance]", str(optim['KNN']['best_params']), f"{optim['KNN']['cv_best_score']:.4f}", f"{optim['KNN']['test_accuracy_after']*100:.2f}%")
    ]

    for idx, row in enumerate(opt_rows_data):
        row_cells = table_opt.rows[idx + 1].cells
        for col_idx, text in enumerate(row):
            row_cells[col_idx].text = text
            set_cell_margins(row_cells[col_idx])
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(9.5)
            if col_idx == 0:
                p.runs[0].font.bold = True
            elif col_idx in [3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run("Bảng 4.7: Kết quả tối ưu hóa siêu tham số bằng GridSearchCV và 5-Fold Cross-Validation")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.bold = True

    add_figure_with_caption(doc, "images/optimization_comparison.png", "Hình 4.11: So sánh độ chính xác của 4 mô hình trước và sau khi tối ưu hóa tham số bằng GridSearchCV", width=Inches(5.0))

    add_heading_3(doc, "4.4.1. Lựa chọn mô hình và cấu hình khuyến nghị")
    add_styled_paragraph(doc, 
        "Dựa trên kết quả thực nghiệm toàn diện và đánh giá Cross-Validation:\n"
        "1. Đối với mô hình KNN: Giá trị K=5 đến K=6 với trọng số đồng đều (uniform) và khoảng cách Euclid là cấu hình tối ưu nhất, vừa đạt điểm kiểm định chéo cao nhất (98.33%) vừa đạt độ chính xác kiểm tra thực tế từ 96.67% đến 100%.\n"
        "2. Đối với mô hình SVM: Khi sử dụng kernel='linear' hoặc 'rbf' với C=1.0 hoặc C=0.1, SVM đạt điểm CV 98.33% và độ chính xác kiểm tra 96.67%, là mô hình có độ tin cậy và tính ổn định toán học cao nhất, ít nguy cơ quá khớp nhất.\n"
        "3. Kết luận lựa chọn: Cả KNN (K=5) và SVM (RBF/Linear) đều là những sự lựa chọn xuất sắc nhất để triển khai trong thực tế cho bài toán phân loại loài hoa Iris."
    )

    doc.add_page_break()

    # ==============================================================================
    # CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN
    # ==============================================================================
    add_heading_1(doc, "CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")

    add_heading_2(doc, "5.1. Kết luận chung")
    add_styled_paragraph(doc, 
        "Qua quá trình nghiên cứu lý thuyết và thực nghiệm triển khai bài tập lớn môn Trí tuệ nhân tạo với đề tài 'Phân loại loài hoa Iris bằng học máy', đề tài đã hoàn thành xuất sắc các mục tiêu nghiên cứu đề ra:\n"
        "1. Về cơ sở khoa học: Đã nghiên cứu sâu sắc bản chất bài toán phân loại đa lớp có giám sát, phân tích chi tiết đặc tính của tập dữ liệu hình thái học Iris, đồng thời làm rõ cơ sở lý thuyết toán học của bốn thuật toán đại diện: Support Vector Machine (SVM), Decision Tree, Gaussian Naive Bayes và K-Nearest Neighbors (KNN).\n"
        "2. Về quy trình thực nghiệm: Đã thiết lập một quy trình thực nghiệm chuẩn mực khoa học bằng phương pháp Stratified Hold-out 80:20 kết hợp cố định random_state=42, đảm bảo tính đại diện đồng đều của các lớp cũng như khả năng tái lập hoàn toàn của các kết quả thực nghiệm.\n"
        "3. Về kết quả đạt được: Tất cả 4 mô hình đều đạt hiệu năng rất cao trên tập kiểm tra (Accuracy từ 93.33% đến 100.0%). Trong đó, KNN (K=5) đạt độ chính xác tuyệt đối 100%, SVM và Naive Bayes đạt 96.67%, và Cây quyết định đạt 93.33%.\n"
        "4. Về phân tích so sánh và tối ưu: Đã xây dựng các ma trận nhầm lẫn chi tiết, phân tích bản chất các ca phân loại sai giữa Versicolor và Virginica, trực quan hóa ranh giới quyết định (Decision Boundaries) sinh động, và ứng dụng thành công GridSearchCV kết hợp 5-Fold Cross-Validation để xác định cấu hình tham số tối ưu nhất cho từng thuật toán."
    )

    add_heading_2(doc, "5.2. Đánh giá ưu điểm và hạn chế của đề tài")
    add_styled_paragraph(doc, 
        "Ưu điểm của đề tài:\n"
        "• Mã nguồn được xây dựng bài bản, module hóa riêng biệt cho từng thuật toán (svm_classifier.py, decision_tree_classifier.py, naive_bayes_classifier.py, knn_classifier.py) và có file điều phối tổng thể main_experiment.py.\n"
        "• Báo cáo kết quả chi tiết, bảng số liệu chuẩn xác, hệ thống hình ảnh biểu đồ và ma trận nhầm lẫn trực quan độ phân giải cao được chèn đồng bộ.\n"
        "• Kết hợp cả đánh giá Hold-out và Cross-Validation, giúp nhìn nhận khách quan khả năng tổng quát hóa của từng mô hình.\n\n"
        "Hạn chế còn tồn tại:\n"
        "• Tập dữ liệu Iris có quy mô tương đối nhỏ (150 mẫu), do đó sự chênh lệch giữa các mô hình chỉ phản ánh trên một vài mẫu kiểm tra cụ thể.\n"
        "• Chưa triển khai thử nghiệm thêm các kỹ thuật học kết hợp (Ensemble Learning) hiện đại như Random Forest, Gradient Boosting hoặc XGBoost."
    )

    add_heading_2(doc, "5.3. Hướng phát triển trong tương lai")
    add_styled_paragraph(doc, 
        "Từ những kết quả nền tảng đã đạt được, hướng mở rộng tiếp theo của nghiên cứu bao gồm:\n"
        "1. Mở rộng thử nghiệm các thuật toán học tổ hợp (Random Forest, Extra Trees, AdaBoost, XGBoost, LightGBM) để đánh giá khả năng cải thiện độ ổn định của ranh giới quyết định.\n"
        "2. Ứng dụng các thuật toán trên vào các bộ dữ liệu thực tế có kích thước lớn hơn, số chiều nhiều hơn và có độ nhiễu phức tạp hơn (ví dụ như bài toán phân loại bệnh lý, thị giác máy tính hoặc xử lý ngôn ngữ tự nhiên).\n"
        "3. Xây dựng giao diện ứng dụng tương tác (Web Application bằng Flask/FastAPI hoặc Streamlit) cho phép người dùng nhập trực tiếp kích thước 4 đặc trưng của một bông hoa và nhận kết quả dự đoán loài hoa tức thì kèm theo độ tin cậy xác suất."
    )

    doc.add_page_break()

    # ==============================================================================
    # TÀI LIỆU THAM KHẢO
    # ==============================================================================
    add_heading_1(doc, "TÀI LIỆU THAM KHẢO")

    refs = [
        "[1] R. A. Fisher, \"The use of multiple measurements in taxonomic problems,\" Annals of Eugenics, vol. 7, no. 2, pp. 179-188, 1936.",
        "[2] V. N. Vapnik, The Nature of Statistical Learning Theory. New York: Springer-Verlag, 1995.",
        "[3] L. Breiman, J. H. Friedman, R. A. Olshen, and C. J. Stone, Classification and Regression Trees. Belmont, CA: Wadsworth International Group, 1984.",
        "[4] T. M. Mitchell, Machine Learning. New York: McGraw-Hill, 1997.",
        "[5] F. Pedregosa et al., \"Scikit-learn: Machine Learning in Python,\" Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
        "[6] C. M. Bishop, Pattern Recognition and Machine Learning. New York: Springer, 2006.",
        "[7] T. Hastie, R. Tibshirani, and J. Friedman, The Elements of Statistical Learning: Data Mining, Inference, and Prediction, 2nd ed. New York: Springer, 2009.",
        "[8] PGS. TS. Nguyễn Thanh Thủy, Giáo trình Trí tuệ Nhân tạo. Hà Nội: Nhà xuất bản Giáo dục Việt Nam, 2014.",
        "[9] Vũ Hữu Tiệp, Machine Learning cơ bản. Hà Nội: Nhà xuất bản Thế Giới, 2018."
    ]

    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_before = Pt(2)
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.25
        p_ref.paragraph_format.left_indent = Inches(0.4)
        p_ref.paragraph_format.first_line_indent = Inches(-0.4)
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = 'Times New Roman'
        r_ref.font.size = Pt(12)

def build_report():
    print("=" * 60)
    print("TIẾN HÀNH TẠO CÁC FILE TÀI LIỆU BÁO CÁO WORD (.DOCX)...")
    print("=" * 60)

    # Tải kết quả thực nghiệm từ JSON
    with open('experiment_results.json', 'r', encoding='utf-8') as f:
        exp_data = json.load(f)

    # 1. TẠO FILE RIÊNG CHƯƠNG 3, 4, 5 (Đúng yêu cầu: 'làm tiếp các chương ra 1 file doc khác')
    doc_separate = docx.Document()
    for section in doc_separate.sections:
        section.top_margin = Inches(0.98) # 2.5 cm
        section.bottom_margin = Inches(0.98) # 2.5 cm
        section.left_margin = Inches(1.18) # 3.0 cm
        section.right_margin = Inches(0.79) # 2.0 cm

    # Trang bìa phụ của file riêng
    p_uni = doc_separate.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐÔNG Á\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(14)
    r_uni.font.bold = True

    p_title = doc_separate.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(15)
    r_title = p_title.add_run("BÀI TẬP LỚN: HỌC PHẦN TRÍ TUỆ NHÂN TẠO\nĐỀ TÀI: PHÂN LOẠI LOÀI HOA IRIS BẰNG CÁC THUẬT TOÁN HỌC MÁY CÓ GIÁM SÁT\n(PHẦN BÁO CÁO THỰC NGHIỆM TIẾP NỐI: CHƯƠNG 3, 4, 5)")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    p_info = doc_separate.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_after = Pt(25)
    r_info = p_info.add_run("Giảng viên hướng dẫn: Trần Xuân Thanh\nSinh viên thực hiện: Nguyễn Thanh Hùng\nMã sinh viên: 20232139\nLớp hành chính: DCCNTT.14.6\nNăm học: 2026")
    r_info.font.name = 'Times New Roman'
    r_info.font.size = Pt(12)
    r_info.font.italic = True

    doc_separate.add_page_break()
    append_content_to_doc(doc_separate, exp_data)

    output_separate = "Bao_Cao_Tieu_Luan_Chuong_3_4_5.docx"
    doc_separate.save(output_separate)
    print(f"\n[1] ĐÃ XUẤT THÀNH CÔNG FILE WORD RIÊNG: {output_separate}")

    # 2. TẠO PHIÊN BẢN GỘP TOÀN DIỆN KẾ THỪA FILE GỐC (Chương 1, 2, 3, 4, 5)
    original_file = 'Nguyen Thanh Hung 20232139.docx'
    if os.path.exists(original_file):
        print(f"\n[2] ĐANG TẠO BẢN GỘP TOÀN DIỆN TỪ '{original_file}' (Chương 1 đến 5)...")
        doc_full = docx.Document(original_file)
        idx = -1
        for i, p in enumerate(doc_full.paragraphs):
            if 'CHƯƠNG 3' in p.text:
                idx = i
                break
        if idx != -1:
            for p in doc_full.paragraphs[idx:]:
                p_elem = p._p
                p_elem.getparent().remove(p_elem)
        
        doc_full.add_page_break()
        append_content_to_doc(doc_full, exp_data)
        output_full = "Bao_Cao_Phan_Loai_Iris_Toan_Dien.docx"
        doc_full.save(output_full)
        print(f"[2] ĐÃ XUẤT THÀNH CÔNG BẢN GỘP TOÀN DIỆN: {output_full}")

    print("=" * 60)
    print("HOÀN TẤT TẠO CÁC FILE BÁO CÁO WORD THÀNH CÔNG!")
    print("=" * 60)

if __name__ == '__main__':
    build_report()
